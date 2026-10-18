"""Builds deck/SotonShip-pitch.pptx: the 10-slide pitch, with notes, transitions and click animations.

Every figure and quote comes from the repo (NUMBERS.md, agent/confirmed_with_websites.csv,
code/messages.py, the track brief). Screenshots in deck/assets/ are of the running site
filled with the 68 synthetic reports.

Run: python deck/build_deck.py   (needs python-pptx and lxml)
"""

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "SotonShip-pitch.pptx"

FONT = "Arial"
MIN_PT = 18

# Palette from code/web/static/kerbside.css.
LIGHT = dict(bg="F4F6F9", surface="FFFFFF", surface2="EEF2F6", ink="0E1726", ink2="3B4859",
             ink3="556273", line="C6D0DB", accent="0369A1", soft="E2EFF8", accent_ink="03507B")
DARK = dict(bg="0A1019", surface="111925", surface2="18222F", ink="E8EEF5", ink2="B6C2D0",
            ink3="95A2B3", line="2A3747", accent="5FB4EC", soft="11314B", accent_ink="5FB4EC")
REVIEW_BG, REVIEW_INK = "FCEFD8", "7F4500"
DONE_BG, DONE_INK = "DFF2E7", "15603B"
HIGH = "B42B31"
BUBBLE_GREY = "E9ECF0"

CHECKED = "checked 17 Oct 2026"


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------- drawing helpers ----------

def text(shapes, x, y, w, h, paras, size=20, color="0E1726", bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, name=None, spacing=1.08, after=0):
    """A text box with no padding. paras: a string (\\n splits paragraphs) or a list of
    paragraphs, each a string or a list of (text, {size, color, bold, italic}) runs."""
    box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    fill_text(box.text_frame, paras, size, color, bold, align, anchor, italic, spacing, after)
    return box


def fill_text(tf, paras, size, color, bold, align, anchor, italic=False, spacing=1.08, after=0,
              margin=0.0):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = paras.split("\n")
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if after:
            p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, opts in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            pt = opts.get("size", size)
            assert pt >= MIN_PT, f"text under {MIN_PT}pt: {run_text!r}"
            f.size = Pt(pt)
            f.name = FONT
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = rgb(opts.get("color", color))


def box(shapes, x, y, w, h, fill, line=None, radius=0.14, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None,
        line_w=1.0):
    s = shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    return s


def label_box(shapes, x, y, w, h, fill, paras, size=20, color="0E1726", bold=False, line=None,
              radius=0.14, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.16,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None, spacing=1.08):
    s = box(shapes, x, y, w, h, fill, line, radius, shape, name)
    fill_text(s.text_frame, paras, size, color, bold, align, anchor, spacing=spacing, margin=margin)
    return s


def circle_num(shapes, x, y, d, n, fill, ink):
    return label_box(shapes, x, y, d, d, fill, str(n), size=20 if d < 0.6 else 24, color=ink, bold=True,
                     align=PP_ALIGN.CENTER, margin=0, shape=MSO_SHAPE.OVAL)


def arrow(shapes, x1, y1, x2, y2, color, width=2.0, head=True, both=False):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if both:
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    if head:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def picture(shapes, path, x, y, w=None, h=None, radius=0.12, line="C6D0DB", name=None):
    pic = shapes.add_picture(str(path), Inches(x), Inches(y),
                             Inches(w) if w else None, Inches(h) if h else None)
    if name:
        pic.name = name
    # Rounded frame: the deck's one motif, shared with every card.
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    short = min(pic.width, pic.height) / 914400
    etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(50000 * min(1, 2 * radius / short))}")
    pic.line.color.rgb = rgb(line)
    pic.line.width = Pt(1)
    return pic


def group(slide, name):
    g = slide.shapes.add_group_shape()
    g.name = name
    return g


def background(slide, hex_):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_)


def header(slide, kicker, title, pal):
    text(slide.shapes, 0.6, 0.42, 9.0, 0.4, kicker.upper(), size=18, bold=True, color=pal["accent"],
         name="!!kicker")
    text(slide.shapes, 0.6, 0.78, 12.1, 0.75, title, size=36, bold=True, color=pal["ink"], name="!!title")


def notes(slide, script, clicks):
    body = script.strip()
    if clicks:
        body += f"\n\nClicks: {clicks} (each click reveals the next part; the slide reads fine with no animation)."
    slide.notes_slide.notes_text_frame.text = body


# ---------- slides ----------

def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, DARK["bg"])
    d = DARK
    s.shapes.add_picture(str(ASSETS / "logo.png"), Inches(0.6), Inches(0.6), Inches(0.9), Inches(0.9))
    text(s.shapes, 0.6, 2.05, 11.0, 0.45, "SotonShip 2026 · Southampton City track", size=22, bold=True,
         color=d["accent"])
    text(s.shapes, 0.6, 2.55, 11.0, 1.3, "Kerbside", size=72, bold=True, color=d["ink"])
    text(s.shapes, 0.6, 3.9, 10.2, 1.2,
         "Text a street problem on iMessage. Get one clear report and a private link back.",
         size=28, color=d["ink2"], spacing=1.12)
    box(s.shapes, 0.6, 5.75, 12.13, 0.02, d["line"], radius=0, shape=MSO_SHAPE.RECTANGLE)
    text(s.shapes, 0.6, 6.0, 12.1, 0.9,
         [[("Team  ", {"bold": True, "color": d["accent"]}),
           ("Krish Mathur · Aadil Sirajudeen · Anirudh Udupi Diggavi · [SID24555: name to confirm]", {})]],
         size=20, color=d["ink2"])
    notes(s, """
Hi, we're the Kerbside team, on the Southampton City track. Kerbside lets a resident text a street problem on iMessage, the app that's already on their phone. It answers only from the council's own web pages, asks a few short questions, and saves one clear report that council staff can act on. The resident gets a private link to follow it. Here's why that matters.
""", 0)
    return s


def slide_problem(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "The problem", "Reporting a problem is harder than it should be", p)
    cards = [
        ("~700", "services run by Southampton City Council.",
         f"About Southampton City Council, {CHECKED}"),
        ("Which council?", "Some SO16, SO18 and SO19 addresses are outside the city. A postcode alone doesn't say.",
         f"Council's missed bin form, {CHECKED}"),
        ("Again and again", "“Residents report fly-tipping, potholes and broken street lights as photos and "
         "free text, often several times for the same problem.”",
         "SotonShip 2026 Southampton City brief"),
    ]
    clicks = []
    for i, (big, body, source) in enumerate(cards):
        x = 0.6 + i * 4.09
        g = group(s, f"Problem card {i + 1}")
        box(g.shapes, x, 1.9, 3.95, 4.75, p["surface"], p["line"])
        text(g.shapes, x + 0.3, 2.15, 3.4, 0.9, big, size=40 if len(big) < 8 else 30, bold=True,
             color=p["accent"], anchor=MSO_ANCHOR.BOTTOM)
        text(g.shapes, x + 0.3, 3.2, 3.4, 2.3, body, size=22, color=p["ink"], spacing=1.1)
        text(g.shapes, x + 0.3, 5.6, 3.4, 0.85, source, size=18, color=p["ink3"], anchor=MSO_ANCHOR.BOTTOM)
        clicks.append([g])
    notes(s, """
Southampton City Council runs around 700 services, by its own count. A resident with a pothole shouldn't need to know which team handles it, or even which council: the council itself says some SO16, SO18 and SO19 addresses are outside the city. And the track brief says residents report the same problem several times, as photos and free text. So the council gets repeated, messy reports.
""", 3)
    return s, clicks


def slide_idea(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "The idea", "Message it like a friend, in the app you already use", p)
    rows = [
        ("Nothing to install", "It's a contact in Messages. Text it like anyone else."),
        ("Words or a photo", "Say what and where. It asks only for what's missing, in five messages at most."),
        ("A private link back", "One report is saved, and the resident gets a link to follow it."),
    ]
    clicks = []
    for i, (head, sub) in enumerate(rows):
        y = 2.0 + i * 1.55
        g = group(s, f"Idea row {i + 1}")
        circle_num(g.shapes, 0.6, y, 0.7, i + 1, p["accent"], "FFFFFF")
        text(g.shapes, 1.6, y - 0.02, 6.4, 0.5, head, size=26, bold=True, color=p["ink"])
        text(g.shapes, 1.6, y + 0.5, 6.4, 0.85, sub, size=20, color=p["ink2"])
        clicks.append([g])

    # Phone mock-up with the agent's real fly-tipping reply (code/messages.py), shortened.
    px, py, pw, ph = 8.75, 1.8, 3.98, 4.6
    box(s.shapes, px, py, pw, ph, p["ink"], radius=0.42)
    box(s.shapes, px + 0.12, py + 0.12, pw - 0.24, ph - 0.24, p["surface"], radius=0.32)
    text(s.shapes, px + 0.2, py + 0.3, pw - 0.4, 0.35, "Kerbside", size=18, bold=True, color=p["ink2"],
         align=PP_ALIGN.CENTER)
    box(s.shapes, px + 0.12, py + 0.78, pw - 0.24, 0.015, p["line"], radius=0, shape=MSO_SHAPE.RECTANGLE)
    label_box(s.shapes, px + 0.55, py + 1.0, pw - 0.8, 0.95, p["accent"],
              "Someone's dumped a mattress on Lawn Road", size=18, color="FFFFFF", radius=0.22)
    label_box(s.shapes, px + 0.22, py + 2.15, pw - 0.44, 1.65, BUBBLE_GREY,
              [[("Logged: fly-tipping, Lawn Road.", {})], [("It hasn't gone to the council yet.", {})],
               [("Follow it ›", {"bold": True, "color": p["accent"]})]],
              size=18, color=p["ink"], radius=0.22, spacing=1.05)
    text(s.shapes, px, py + ph + 0.1, pw, 0.4, "Real reply text, shortened", size=18,
         color=p["ink3"], align=PP_ALIGN.CENTER)
    notes(s, """
Our idea: message it like a friend, in the app you already use. There's nothing to install: it's a contact in Messages. You say what and where, in words or a photo, and it asks only for what's missing, five messages at most. Then it saves one report and sends you a private link to follow it. The reply on the right is the agent's real wording.
""", 3)
    return s, clicks


def slide_how(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "How it works", "One message in, one report out", p)
    W, H = 2.75, 1.95
    X = [0.6, 3.727, 6.853, 9.98]
    Y1, Y2 = 1.8, 4.3
    steps = [
        (X[0], Y1, "Resident", "texts on iMessage"),
        (X[1], Y1, "Agent", "asks up to five follow-ups"),
        (X[1], Y2, "Verified facts", "54, each with a source and date"),
        (X[2], Y1, "Reports", "one per problem, repeats checked"),
        (X[3], Y1, "Dashboard", "council staff filter and update"),
        (X[2], Y2, "Private link", "resident sees only their report"),
    ]
    line = p["ink3"]
    clicks = []
    for i, (x, y, head, sub) in enumerate(steps):
        g = group(s, f"Step {i + 1}")
        agent = head == "Agent"
        fill, ink, ink2 = (p["accent"], "FFFFFF", "FFFFFF") if agent else (p["surface"], p["ink"], p["ink2"])
        box(g.shapes, x, y, W, H, fill, None if agent else p["line"])
        circle_num(g.shapes, x + 0.18, y + 0.18, 0.46, i + 1, "FFFFFF" if agent else p["soft"],
                   p["accent"] if agent else p["accent_ink"])
        text(g.shapes, x + 0.2, y + 0.72, W - 0.4, 0.45, head, size=22, bold=True, color=ink)
        text(g.shapes, x + 0.2, y + 1.2, W - 0.4, 0.7, sub, size=18, color=ink2)
        mid1, mid2 = Y1 + H / 2, Y2 + H / 2
        if i == 1:      # resident -> agent
            arrow(g.shapes, X[0] + W + 0.05, mid1, X[1] - 0.05, mid1, line)
        elif i == 2:    # agent <-> facts
            arrow(g.shapes, X[1] + W / 2, Y1 + H + 0.05, X[1] + W / 2, Y2 - 0.05, line, both=True)
        elif i == 3:    # agent -> database
            arrow(g.shapes, X[1] + W + 0.05, mid1, X[2] - 0.05, mid1, line)
        elif i == 4:    # database -> dashboard
            arrow(g.shapes, X[2] + W + 0.05, mid1, X[3] - 0.05, mid1, line)
        elif i == 5:    # database -> link, then the link back to the resident along the bottom
            arrow(g.shapes, X[2] + W / 2, Y1 + H + 0.05, X[2] + W / 2, Y2 - 0.05, line)
            bx, by = X[2] + W / 2, Y2 + H
            low = 6.7
            arrow(g.shapes, bx, by + 0.05, bx, low, p["accent"], head=False)
            arrow(g.shapes, bx, low, X[0] + W / 2, low, p["accent"], head=False)
            arrow(g.shapes, X[0] + W / 2, low, X[0] + W / 2, Y1 + H + 0.05, p["accent"])
            text(g.shapes, X[0] + W / 2 + 0.2, low - 0.45, 4.6, 0.4, "link sent back in the chat", size=18,
                 color=p["accent_ink"], bold=True)
        clicks.append([g])
    notes(s, """
Here's how it works. A resident texts on iMessage. The agent reads it, asks up to five short follow-ups, and checks 54 verified facts, each with a council source and a date. It saves one report to the database and checks for repeats. Staff see it on the council dashboard. The resident gets a private link back in the chat, showing only their own report.
""", 6)
    return s, clicks


def slide_trust(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "Trust", "It only says what the council's own pages say", p)
    facts = [
        ("Over 7.5cm", "Road potholes over 7.5cm: repaired within 24 hours.", "Council potholes page"),
        ("75%", "Aim: 75% of urgent street cleaning requests dealt with in one working day. An aim, not a "
         "guarantee.", "Environmental issues page"),
        ("023 8023 3344", "Council housing emergency repairs, out of office hours.", "Out of hours page"),
    ]
    clicks = []
    for i, (big, body, source) in enumerate(facts):
        x = 0.6 + i * 4.09
        g = group(s, f"Fact {i + 1}")
        box(g.shapes, x, 1.85, 3.95, 3.5, p["surface"], p["line"])
        text(g.shapes, x + 0.3, 2.05, 3.4, 0.7, big, size=34 if len(big) < 12 else 30, bold=True,
             color=p["accent"], anchor=MSO_ANCHOR.MIDDLE)
        text(g.shapes, x + 0.3, 2.85, 3.4, 1.55, body, size=20, color=p["ink"])
        text(g.shapes, x + 0.3, 4.45, 3.4, 0.75, [source, CHECKED], size=18, color=p["ink3"],
             anchor=MSO_ANCHOR.BOTTOM)
        clicks.append([g])
    g = group(s, "Cannot confirm")
    box(g.shapes, 0.6, 5.55, 12.13, 1.35, p["soft"])
    text(g.shapes, 0.9, 5.55, 3.1, 1.35, "When the facts don't cover it", size=22, bold=True,
         color=p["accent_ink"], anchor=MSO_ANCHOR.MIDDLE)
    text(g.shapes, 4.2, 5.55, 8.3, 1.35,
         "“I can't confirm who owns that road. Try the council's roads page or National Highways.”",
         size=22, italic=True, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE)
    clicks.append([g])
    notes(s, """
Trust is the core rule. The agent only states what the council's own pages say, and shows the source and the date we checked. Potholes over 7.5 centimetres: within 24 hours. The 75 percent street cleaning aim, which the council calls an aim, not a guarantee. The out-of-hours number. And when the facts don't cover it, it says so, in exactly these words.
""", 4)
    return s, clicks


def slide_demo_resident(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "Demo · resident", "A conversation, then a private link", p)
    cx, cy, cw, ch = 0.6, 1.8, 6.3, 4.95
    box(s.shapes, cx, cy, cw, ch, p["surface"], p["line"])
    text(s.shapes, cx, cy + 0.15, cw, 0.35, "Messages · Kerbside", size=18, bold=True, color=p["ink2"],
         align=PP_ALIGN.CENTER)
    box(s.shapes, cx, cy + 0.6, cw, 0.015, p["line"], radius=0, shape=MSO_SHAPE.RECTANGLE)
    right = cx + cw - 0.2
    turns = [
        ("you", 4.6, 0.55, "Massive pothole, nearly burst a tyre"),
        ("bot", 5.6, 1.05, ["Logged: pothole.", "Which road, and nearest house number or junction?"]),
        ("you", 3.3, 0.55, "46 Portswood Road"),
        ("bot", 5.6, 1.15, ["Thanks, added: pothole at Portswood Road.", "It hasn't gone to the council yet.",
                            [("Follow it ›", {"bold": True, "color": p["accent"]})]]),
    ]
    y = cy + 0.8
    clicks = []
    g = None
    for who, w, h, body in turns:
        if who == "you":
            g = group(s, f"Chat turn {len(clicks) + 1}")
            label_box(g.shapes, right - w, y, w, h, p["accent"], body, size=18, color="FFFFFF", radius=0.2)
        else:
            label_box(g.shapes, cx + 0.2, y, w, h, BUBBLE_GREY, body, size=18, color=p["ink"], radius=0.2,
                      spacing=1.05)
            clicks.append([g])
        y += h + 0.14
    text(s.shapes, cx, cy + ch + 0.1, cw, 0.4, "The agent's real reply text, shortened", size=18,
         color=p["ink3"], align=PP_ALIGN.CENTER)

    g = group(s, "Resident page")
    text(g.shapes, 7.3, 2.0, 2.45, 0.5, "Tap the link", size=24, bold=True, color=p["accent"])
    text(g.shapes, 7.3, 2.6, 2.45, 2.0, "Status, next step and progress, for this one report.", size=20,
         color=p["ink"], spacing=1.1)
    text(g.shapes, 7.3, 4.25, 2.45, 1.4, "No report ID or name in the link.", size=20, color=p["ink2"],
         spacing=1.1)
    label_box(g.shapes, 7.3, 5.75, 2.45, 0.95, REVIEW_BG, ["Real page,", "synthetic data"], size=18,
              color=REVIEW_INK, bold=True, radius=0.16, align=PP_ALIGN.CENTER)
    picture(g.shapes, ASSETS / "resident-page.jpg", 10.1, 1.8, h=5.25, radius=0.18)
    clicks.append([g])
    notes(s, """
Now the demo, from the resident's side. They text: massive pothole. The agent logs it and asks one thing: which road. They answer, and it confirms, says honestly that it hasn't gone to the council yet, and sends a link. The link opens this page: status, what happens next, and progress. It's a real page from our site; the report on it is synthetic.
""", 3)
    return s, clicks


def slide_demo_council(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "Demo · council", "Every report on one dashboard", p)
    label_box(s.shapes, 9.35, 0.85, 3.38, 0.55, REVIEW_BG, "68 synthetic reports", size=18, color=REVIEW_INK,
              bold=True, radius=0.27, align=PP_ALIGN.CENTER)
    dw, dh = 7.04, 4.4
    pw = 4.4 * 760 / 1079
    x0 = (13.333 - (dw + 0.4 + pw)) / 2
    picture(s.shapes, ASSETS / "dashboard-list.jpg", x0, 1.75, w=dw, h=dh, name="Dashboard desktop")
    g1 = group(s, "Dashboard phone")
    picture(g1.shapes, ASSETS / "dashboard-phone.jpg", x0 + dw + 0.4, 1.75, w=pw, h=dh, radius=0.18)
    g2 = group(s, "Dashboard features")
    chips = ["Filters and search", "Priority + reasons", "Status in one click", "Repeats joined up"]
    for i, chip in enumerate(chips):
        label_box(g2.shapes, 0.6 + i * 3.08, 6.4, 2.93, 0.6, p["surface"], chip, size=18, color=p["accent_ink"],
                  bold=True, line=p["line"], radius=0.3, align=PP_ALIGN.CENTER, margin=0.08)
    notes(s, """
And the council side. Every saved report lands on one dashboard. Staff filter by time, category and severity, sort by priority with the reasons shown, and change status in one click. Repeats are joined to the earlier report. It works at phone width too. These are real screenshots of our running site, filled with 68 synthetic reports.
""", 2)
    return s, [[g1], [g2]]


def slide_safety(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "Safety and privacy", "Knows when to stop. Keeps almost nothing.", p)
    cols = [
        ("Points elsewhere, saves nothing", [
            ("Emergency", "“Call 999 now.”"),
            ("Outside the city", "their own council"),
            ("Not council", "National Highways, utility"),
            ("Neighbour issues", "the council's own form"),
        ]),
        ("What a report keeps", [
            ("Who", "a salted hash, not the number"),
            ("Link", "random 16-character token"),
            ("Where", "the street, no coordinates"),
            ("Dashboard", "staff only, password"),
        ]),
    ]
    clicks = []
    for i, (head, rows) in enumerate(cols):
        x = 0.6 + i * 6.18
        g = group(s, f"Safety column {i + 1}")
        box(g.shapes, x, 1.8, 5.95, 3.95, p["surface"], p["line"])
        text(g.shapes, x + 0.3, 1.98, 5.4, 0.5, head, size=24, bold=True, color=p["accent_ink"])
        for j, (k, v) in enumerate(rows):
            text(g.shapes, x + 0.3, 2.7 + j * 0.74, 5.45, 0.7,
                 [[(k + "  ", {"bold": True}), (v, {"color": p["ink2"]})]], size=20, color=p["ink"])
        clicks.append([g])
    g = group(s, "Credits")
    box(g.shapes, 0.6, 5.95, 12.13, 1.0, p["surface2"])
    text(g.shapes, 0.85, 5.95, 11.65, 1.0,
         [[("Credits  ", {"bold": True, "color": p["ink"]}),
           ("Facts: Southampton City Council web pages, checked 17 Oct 2026 · Testing only: FixMyStreet reports "
            "(mySociety), via Open311 · Streets and city boundary: © OpenStreetMap contributors, ODbL", {})]],
         size=18, color=p["ink2"], anchor=MSO_ANCHOR.MIDDLE)
    clicks.append([g])
    notes(s, """
It knows when to stop. An emergency gets: call 999 now. A place outside the city gets pointed to its own council, using the real boundary, not a postcode guess. Trunk roads, utilities and neighbour issues go to the right place. None of those are saved. A report keeps a salted hash, never the phone number, and a random private link. Credits are on screen.
""", 3)
    return s, clicks


def slide_results(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    p = LIGHT
    background(s, p["bg"])
    header(s, "Results", "What we measured: wins and losses", p)
    stats = [
        ("42%", "Keyword baseline: 20 of 48 categories right", p["accent"]),
        ("Not measured yet", "Our agent's category accuracy", p["ink3"]),
        ("8 of 10", "real repeats caught by the duplicate rule", p["accent"]),
        ("17 of 40", "non-repeats wrongly flagged as repeats", HIGH),
    ]
    clicks = [[], []]
    for i, (big, label, color) in enumerate(stats):
        x = 0.6 + i * 3.09
        g = group(s, f"Stat {i + 1}")
        box(g.shapes, x, 1.8, 2.85, 2.35, p["surface"], p["line"])
        text(g.shapes, x + 0.25, 1.95, 2.4, 0.85, big, size=36 if len(big) < 10 else 24, bold=True, color=color,
             anchor=MSO_ANCHOR.MIDDLE)
        text(g.shapes, x + 0.25, 2.9, 2.4, 1.15, label, size=20, color=p["ink"])
        clicks[i // 2].append(g)
    lists = [
        ("Wins", DONE_INK, ["Fly-tipping: 6 of 6 right", "Exact double submissions: 2 of 2",
                            "Look-alike pairs: 0 of 19 flagged"]),
        ("Where it fails", HIGH, ["Services we don't handle: 0 of 17", "Litter called fly-tipping: 4 of 6",
                                  "Busy streets: new problems merged"]),
    ]
    for i, (head, color, items) in enumerate(lists):
        x = 0.6 + i * 6.18
        g = group(s, head)
        box(g.shapes, x, 4.35, 5.95, 1.95, p["surface"], p["line"])
        box(g.shapes, x, 4.35, 0.1, 1.95, color, radius=0, shape=MSO_SHAPE.RECTANGLE)
        text(g.shapes, x + 0.35, 4.45, 5.4, 0.45, head, size=22, bold=True, color=color)
        text(g.shapes, x + 0.35, 4.95, 5.4, 1.3, items, size=20, color=p["ink"])
        clicks.append([g])
    text(s.shapes, 0.6, 6.42, 12.13, 0.7,
         "Source: NUMBERS.md. 60 real Southampton FixMyStreet reports and 70 pairs, labelled by us; ambiguous "
         "cases left out. Fetched 17 Oct 2026.", size=18, color=p["ink3"])
    notes(s, """
What we measured, honestly. On 60 real Southampton FixMyStreet reports, a keyword baseline got 20 of 48 right. We haven't measured our agent yet. The duplicate rule caught 8 of 10 real repeats, but wrongly flagged 17 of 40, all new problems on busy streets. Wins on the left, failures on the right. It's a small sample, labelled by us.
""", 4)
    return s, clicks


def slide_next(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    d = DARK
    background(s, d["bg"])
    header(s, "What's next", "From a demo to a council pilot", d)
    rows = [
        ("Measure the agent", "Accuracy and cost per report, on the same test set."),
        ("Send reports to the council", "Today they are saved here only. Next: the council's own forms."),
        ("Try it with one council team", "“Every UK council handles the same kind of reports.” (track brief)"),
    ]
    clicks = []
    for i, (head, sub) in enumerate(rows):
        y = 1.95 + i * 1.55
        g = group(s, f"Next {i + 1}")
        circle_num(g.shapes, 0.6, y, 0.7, i + 1, d["accent"], d["bg"])
        text(g.shapes, 1.6, y - 0.02, 6.6, 0.5, head, size=26, bold=True, color=d["ink"])
        text(g.shapes, 1.6, y + 0.5, 6.6, 0.85, sub, size=20, color=d["ink2"])
        clicks.append([g])
    box(s.shapes, 8.65, 1.85, 4.08, 4.6, d["surface"], d["line"])
    text(s.shapes, 8.95, 2.1, 3.5, 0.5, "Team", size=24, bold=True, color=d["accent"])
    text(s.shapes, 8.95, 2.75, 3.5, 2.6,
         ["Krish Mathur", "Aadil Sirajudeen", "Anirudh Udupi Diggavi", "[SID24555: name to confirm]"],
         size=20, color=d["ink"], after=8)
    text(s.shapes, 8.95, 5.45, 3.5, 0.8, "SotonShip 2026\nSouthampton City track", size=18, color=d["ink3"])
    notes(s, """
Next: measure the agent, and its cost per report, on the same test set. Send reports into the council's own forms; today they're saved here only. Then try it with one council team. The brief says every UK council handles the same kind of reports, so this can travel. Thank you. We're happy to take questions.
""", 3)
    return s, clicks


# ---------- transitions and animations (raw PresentationML) ----------

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"

FADE = f'<p:transition xmlns:p="{P}" spd="med"><p:fade/></p:transition>'
MORPH = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{P}">'
         f'<mc:Choice xmlns:p159="{P159}" Requires="p159">'
         f'<p:transition spd="slow"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         f'<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback>'
         f'</mc:AlternateContent>')


def timing_xml(clicks, dur=400):
    """Click-by-click fade entrances. clicks: a list of lists of shape ids; the first id in each
    list starts on the click, the rest start with it."""
    n = [2]

    def nid():
        n[0] += 1
        return n[0]

    parts = []
    for spids in clicks:
        outer, inner = nid(), nid()
        effects = []
        for k, spid in enumerate(spids):
            eff, setid, animid = nid(), nid(), nid()
            node = "clickEffect" if k == 0 else "withEffect"
            effects.append(
                f'<p:par><p:cTn id="{eff}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" '
                f'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{setid}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
                f'</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst>'
                f'<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to>'
                f'<p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{animid}" dur="{dur}"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>')
        parts.append(
            f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            f'<p:childTnLst><p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/>'
            f'</p:stCondLst><p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par>'
            f'</p:childTnLst></p:cTn></p:par>')
    return (
        f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
        f'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" '
        f'nodeType="mainSeq"><p:childTnLst>{"".join(parts)}</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


def add_motion(slide, clicks, morph=False):
    sld = slide._element
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    trans = etree.fromstring(MORPH if morph else FADE)
    anchor.addnext(trans)
    if clicks:
        ids = [[g.shape_id for g in step] for step in clicks if step]
        trans.addnext(etree.fromstring(timing_xml(ids)))


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Kerbside: SotonShip 2026 pitch"
    prs.core_properties.author = "Kerbside team"

    built = [(slide_title(prs), [])]
    for make in (slide_problem, slide_idea, slide_how, slide_trust, slide_demo_resident,
                 slide_demo_council, slide_safety, slide_results, slide_next):
        built.append(make(prs))
    for number, (slide, clicks) in enumerate(built, start=1):
        # One transition for the deck; Morph once, where the talk turns from the idea to the demo.
        add_motion(slide, clicks, morph=(number == 6))
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
